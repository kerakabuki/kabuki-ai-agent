# -*- coding: utf-8 -*-
"""川上衣装への見積依頼（FAX用）Word文書を生成する。
実行: python scripts/content/make_kawakami_fax.py [出力パス]
内容の正本は docs/pr/kawakami-isho-mitsumori-request.md。
"""
import sys, os, datetime
from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

OUT = sys.argv[1] if len(sys.argv) > 1 else r"C:\Users\NAO\OneDrive\Documents\気良歌舞伎事務\output\川上衣装_見積依頼_縄山茶屋_2026-10.docx"
TODAY = datetime.date.today()
DATE_JP = f"令和8年{TODAY.month}月{TODAY.day}日"

doc = Document()
for s in doc.sections:
    s.page_height, s.page_width = Mm(297), Mm(210)
    s.left_margin = s.right_margin = Mm(20); s.top_margin = s.bottom_margin = Mm(14)

st = doc.styles["Normal"]; st.font.name = "游明朝"; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "游明朝")

def para(text="", size=None, bold=False, align=None, space_after=4):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = bold
    if size: r.font.size = Pt(size)
    if align == "c": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if align == "r": p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(space_after)
    return p

def heading(text, size=12):
    p = para(text, size=size, bold=True, space_after=3); p.paragraph_format.space_before = Pt(8); return p

def table(rows, widths=None, header=True, size=9.5):
    t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            c = t.cell(i, j); c.text = ""; p = c.paragraphs[0]; r = p.add_run(cell); r.font.size = Pt(size)
            if header and i == 0: r.bold = True
            if widths: c.width = Mm(widths[j])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

# ===== 1枚目: FAX送付状 =====
para("FAX送信票", size=16, bold=True, align="c", space_after=10)
table([
    ["送信先", "川上衣装　御中　　FAX：053-231-2543"],
    ["送信元", "気良歌舞伎保存会 事務局　澤 奈央也　　TEL：090-7040-9899　FAX：0575-87-2435（気良歌舞伎保存会 清水祐作方）"],
    ["送信日", DATE_JP],
    ["枚数", "本票を含め 2 枚"],
    ["件名", "出張公演（岐阜県大垣市・縄山茶屋）向け 衣裳・かつら・化粧の見積のお願い"],
], widths=[30, 140], header=False, size=10.5)
para("")
para("いつもお世話になっております。気良歌舞伎保存会事務局の澤奈央也です。", space_after=6)
para("このたび、岐阜県大垣市の縄山茶屋株式会社より、海外からのお客様向けの出張公演についての見積依頼がありました。"
     "つきましては、下記の三つの形について、衣裳・かつら・化粧の見積をお願いしたく、別紙のとおり役ごとの内容をお送りします。", space_after=6)
para("会場が大垣市となり、出張費が郡上とは異なるかと存じますので、会場の所在地を明記しております。"
     "お見積りのご返信は上記FAX（清水方）へお願いいたします。ご不明な点はお電話でご連絡ください。何卒よろしくお願い申し上げます。", space_after=10)

heading("1. 会場と条件")
table([
    ["会場", "縄山茶屋株式会社　施設（離れの座敷・地下ホール・屋外舞台）"],
    ["所在地", "〒503-0934 岐阜県大垣市外渕二丁目165番地1（名神高速 大垣ICから約3分。駐車場あり）"],
    ["日程", "土日祝日。3か月前までに日程を確定します。年に2回程度を予定"],
    ["時間帯", "昼の開催か夜の開催かは現時点では未定です。稲瀬川は屋外の舞台で下駄で演じます（雨天は地下ホール）"],
    ["見積の宛先", "企画段階ですので、見積は気良歌舞伎宛で結構です"],
], widths=[30, 140], header=False, size=10)

heading("2. 見積をお願いする三つの形")
table([
    ["", "内容", "衣裳", "化粧", "着替え"],
    ["①", "浜松屋（短縮版）＋稲瀬川（地下ホールで浜松屋 40分前後の見込み、30分の幕間、屋外で稲瀬川 約20分＋撮影タイム）", "14組", "7人", "幕間に7人全員"],
    ["②", "稲瀬川のみ（屋外。約20分＋撮影タイム。捕手2人）", "7組", "5人", "なし"],
    ["③", "浜松屋（短縮版）のみ（地下ホール）", "7組", "7人", "弁天の肌脱ぎのみ（舞台上）"],
], widths=[8, 100, 16, 16, 30])

doc.add_page_break()

# ===== 別紙 =====
para("別紙　役ごとの衣裳・かつら・化粧", size=13, bold=True, align="c", space_after=8)

heading("① 浜松屋（短縮版・40分前後の見込み）＋稲瀬川（兼ね役で7人）　※②と③の合計（下記）。幕間30分で7人全員が着替えます")
table([
    ["役者", "浜松屋（短縮版）", "稲瀬川", "幕間にすること"],
    ["A", "弁天（娘）", "弁天（男）", "衣装替え・化粧直し"],
    ["B", "南郷（侍）", "南郷", "衣装替え・化粧直し"],
    ["C", "逸当（侍）", "駄右衛門", "衣装替え・化粧直し"],
    ["D", "宗之助（若旦那）", "赤星", "衣装替え・化粧直し"],
    ["E", "鳶頭", "忠信", "衣装替え・化粧直し"],
    ["F", "番頭", "捕手", "衣装替え・化粧直し"],
    ["G", "手代", "捕手", "衣装替え・化粧直し"],
], widths=[14, 36, 30, 90])
para("衣裳14組（浜松屋7、揃い5、捕手2）、肉襦袢1。", size=9.5)

heading("② 稲瀬川のみ（5役＋捕手2）　※約20分、終演後にお客様との撮影タイム")
para("役: 五人男の揃いで、弁天小僧菊之助／忠信利平／赤星十三郎／南郷力丸／日本駄右衛門。ほかに捕手2人", size=10)
para("傘5本と十手は気良歌舞伎で用意します。屋外は下駄で演じます（雨天は地下ホール）。", size=9.5)

heading("③ 浜松屋（短縮版）のみ（7役）　※浜松屋見世先の場の短縮版（40分前後の見込み）。幸兵衛・丁稚・若党・中間は出しません")
para("役: 弁天小僧菊之助（屋敷娘お浪。肌脱ぎで桜の彫り物を見せます）／南郷力丸（供の侍）／浜松屋倅 宗之助（若旦那）／番頭 与九郎／手代 佐兵衛／鳶頭 清次／玉島逸当（実は日本駄右衛門。侍の拵え）", size=10)
para("小道具（反物、算盤、燭台、胴巻、金包み）は気良で用意します。", size=9.5)

os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc.save(OUT)
print("saved:", OUT)
